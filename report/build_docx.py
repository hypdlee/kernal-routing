"""
Word version of the two-backbone paper, generated from the same built LaTeX
(report/out/paper2/main.tex and main.aux) as the PDF, so the two can never
disagree. Run build_paper2.py first, then:

    python report/build_docx.py [--src report/out/paper2] [--out paper/two_backbones/KernelRoute_paper.docx]

Every value comes from the generated macros, every section, figure, table
and equation number from main.aux (the numbers LaTeX assigned), every
citation from the bibliography in the template. The script stops if any
macro, reference or citation is left unresolved.
"""
import argparse, pathlib, re, shutil, subprocess, sys, tempfile

HERE = pathlib.Path(__file__).resolve().parent


def arg_end(s, i):
    """s[i] == '{'; return index just past the matching '}'."""
    assert s[i] == "{", s[i:i + 30]
    depth = 0
    for j in range(i, len(s)):
        c = s[j]
        if c == "\\":
            continue
        if c == "{" and s[j - 1] != "\\":
            depth += 1
        elif c == "}" and s[j - 1] != "\\":
            depth -= 1
            if depth == 0:
                return j + 1
    raise ValueError("unbalanced braces at " + s[i:i + 60])


def replace_cmd(s, name, fn, nargs=1):
    """Replace \\name{a}{b}... (nargs braced arguments) by fn(*args)."""
    out, i, pat = [], 0, "\\" + name + "{"
    while True:
        k = s.find(pat, i)
        if k < 0:
            out.append(s[i:]); return "".join(out)
        nxt = s[k + len(pat) - 1:k + len(pat)]
        args, j = [], k + len(pat) - 1
        for _ in range(nargs):
            while s[j] in " \n":
                j += 1
            e = arg_end(s, j); args.append(s[j + 1:e - 1]); j = e
        out.append(s[i:k]); out.append(fn(*args)); i = j


def build(src, out):
    src = pathlib.Path(src); out = pathlib.Path(out)
    tex = (src / "main.tex").read_text()
    aux = (src / "main.aux").read_text()

    # values and labels -----------------------------------------------------
    macros = {}
    for m in re.finditer(r"\\expandafter\\def\\csname v(\w+)\\endcsname\{", tex):
        e = arg_end(tex, m.end() - 1)          # values such as $<10^{-4}$ contain braces
        macros[m.group(1)] = tex[m.end():e - 1]
    labels = {}
    for m in re.finditer(r"\\newlabel\{([^}]+)\}\{\{([^}]*)\}", aux):
        labels[m.group(1)] = m.group(2)
    bib = {}
    for m in re.finditer(r"\\bibitem\[([^\]]+)\]\{(\w+)\}", tex):
        lab = m.group(1).replace("~", " ")
        a, y = re.match(r"(.*)\((\d{4})\)", lab).groups()
        bib[m.group(2)] = (a.strip(), y)

    title_block = tex[tex.index("\\title{"):tex.index("\\begin{document}")]
    body = tex[tex.index("\\begin{abstract}"):tex.index("\\end{document}")]

    # title, author, date, footnote
    t_arg = title_block[title_block.index("\\title{") + 6:]
    t_arg = t_arg[:arg_end(t_arg, 0)][1:-1]
    thanks = replace_cmd(t_arg, "thanks", lambda x: "")
    note = re.search(r"\\thanks\{(.*?)\}\}?$", t_arg, re.S)
    note = t_arg[t_arg.index("\\thanks{") + 7:]
    note = note[1:arg_end(note, 0) - 1]
    note = re.sub(r"\\V\{(\w+)\}", lambda m: macros[m.group(1)], note)
    title_main = re.sub(r"\\vspace\{[^}]*\}|\\textbf|\\large|\\\\\[3pt\]", "", thanks)
    title_main = title_main.replace("{", "").replace("}", "").strip()
    first, sub = [x.strip() for x in title_main.split("\n", 1)] if "\n" in title_main else (title_main, "")
    author = re.search(r"\\author\{Ankit Kumar\\\\\\small (.*?)\}", title_block).group(1)
    date = re.search(r"\\date\{\\small (.*?)\}", title_block).group(1)

    s = body
    # \V values (repeat: a value can never contain \V, but be safe)
    for _ in range(2):
        s = re.sub(r"\\V\{(\w+)\}", lambda m: macros[m.group(1)], s)
    left = re.findall(r"\\V\{\w+\}", s)
    assert not left, left

    # proposition and proof
    s = re.sub(r"\\begin\{proposition\}(?:\[\{(.*?)\}\])?\\label\{prop:collapse\}",
               lambda m: "\\textbf{Proposition " + labels["prop:collapse"] + "}" + (" (" + m.group(1) + ")" if m.group(1) else "") + ". \\emph{", s)
    s = s.replace("\\end{proposition}", "}\n")
    s = s.replace("\\begin{proof}", "\\emph{Proof.} ").replace("\\end{proof}", " $\\square$\n")

    # numbered equations: every equation environment, in order, checked against main.aux
    eqs = list(re.finditer(r"\\begin\{equation\}(.*?)\\end\{equation\}", s, re.S))
    for n, m in reversed(list(enumerate(eqs, 1))):
        body = m.group(1)
        lab = re.search(r"\\label\{([^}]+)\}", body)
        if lab:
            assert labels[lab.group(1)] == str(n), (lab.group(1), labels[lab.group(1)], n)
            body = body.replace(lab.group(0), "")
        s = s[:m.start()] + "$$" + body.strip() + "\\qquad(" + str(n) + ")$$" + s[m.end():]
    assert "\\begin{equation}" not in s
    # tables: plain arm names, unbreakable dataset names, text dagger (narrow columns wrap otherwise)
    def tidy_table(m):
        b = m.group(0)
        # math that only joins arm names ($\texttt{gap}-\texttt{static}$) becomes text
        b = re.sub(r"\$((?:\\texttt\{[^}]*\}|[-\s])+)\$",
                   lambda m: m.group(1).replace("-", " \u2212 ").replace("  ", " "), b)
        b = re.sub(r"\\texttt\{([^}]*)\}", r"\1", b)
        b = b.replace("CIFAR-100", "CIFAR\u2011100").replace("Tiny-ImageNet", "Tiny\u2011ImageNet")
        return b
    s = re.sub(r"\\begin\{tabular\}.*?\\end\{tabular\}", tidy_table, s, flags=re.S)
    s = s.replace("$^\\dagger$", "\\textsuperscript{\u2020}")
    # table rules pandoc cannot read
    s = re.sub(r"\\cmidrule(\([a-z]*\))?\{[^}]*\}", "", s)
    # p < 10^-4 as plain text, so every viewer renders the sign
    s = s.replace("$<10^{-4}$", "< 10\\textsuperscript{\u22124}")

    # captions: prefix with the number LaTeX gave the float
    def number_captions(s):
        out, i = [], 0
        for m in re.finditer(r"\\begin\{(figure|table)\}", s):
            pass
        pos = 0
        res = []
        while True:
            m = re.compile(r"\\begin\{(figure|table)\}").search(s, pos)
            if not m:
                res.append(s[pos:]); break
            env = m.group(1)
            end = s.index("\\end{" + env + "}", m.end())
            block = s[m.start():end]
            lab = re.search(r"\\label\{([^}]+)\}", block).group(1)
            num = labels[lab]
            word = "Figure" if env == "figure" else "Table"
            block = block.replace("\\caption{", "\\caption{" + word + " " + num + ": ", 1)
            res.append(s[pos:m.start()]); res.append(block); pos = end
        return "".join(res)
    s = number_captions(s)

    # references and citations
    s = re.sub(r"\\eqref\{([^}]+)\}", lambda m: "(" + labels[m.group(1)] + ")", s)
    s = re.sub(r"\\ref\{([^}]+)\}", lambda m: labels[m.group(1)], s)

    def cite(keys, paren):
        ks = [k.strip() for k in keys.split(",")]
        if paren:
            parts, i = [], 0
            while i < len(ks):  # group consecutive same-author keys: Ding et al., 2019, 2021
                a, y = bib[ks[i]]; ys = [y]; j = i + 1
                while j < len(ks) and bib[ks[j]][0] == a:
                    ys.append(bib[ks[j]][1]); j += 1
                parts.append(f"{a}, {', '.join(ys)}"); i = j
            return "(" + "; ".join(parts) + ")"
        return "; ".join(f"{bib[k][0]} ({bib[k][1]})" for k in ks)
    s = re.sub(r"\\citep\[([^\]]+)\]\{([^}]+)\}", lambda m: cite(m.group(2), True)[:-1] + ", " + m.group(1) + ")", s)
    s = re.sub(r"\\citealp\[([^\]]+)\]\{([^}]+)\}", lambda m: cite(m.group(2), True)[1:-1] + ", " + m.group(1), s)
    s = re.sub(r"\\citealp\{([^}]+)\}", lambda m: cite(m.group(1), True)[1:-1], s)
    s = re.sub(r"\\citep\{([^}]+)\}", lambda m: cite(m.group(1), True), s)
    s = re.sub(r"\\citet\{([^}]+)\}", lambda m: cite(m.group(1), False), s)
    assert "\\cite" not in s

    # units
    s = re.sub(r"\\SI\{([^}]*)\}\{pp\}", r"\1 pp", s)
    s = re.sub(r"\\SI\{([^}]*)\}\{\\percent\}", r"\1\\%", s)
    s = re.sub(r"\\SI\{([^}]*)\}\{\\giga\\byte\}", r"\1 GB", s)
    s = re.sub(r"\\SI\{([^}]*)\}\{M\}", r"\1 M", s)
    s = s.replace("\\si{\\percent}", "\\%")
    assert "\\SI{" not in s and "\\si{" not in s, re.findall(r"\\S?[iI]\{[^}]*\}\{[^}]*\}", s)[:5]

    # wrappers and layout commands
    s = s.replace("\\resizebox{\\textwidth}{!}{%\n", "")
    s = re.sub(r"\\end\{tabular\}\}", r"\\end{tabular}", s)
    for c in ["\\FloatBarrier", "\\clearpage", "\\centering", "\\small", "\\scriptsize", "\\noindent"]:
        s = s.replace(c, "")
    s = re.sub(r"\\setlength\\itemsep\{[^}]*\}", "", s)
    s = re.sub(r"\[(!htbp|h)\]", "", s)
    s = s.replace("\\zv", "\\mathbf{z}")
    s = re.sub(r"\\label\{[^}]+\}", "", s)

    # section numbers as LaTeX assigned them
    app = s.index("\\appendix")
    nsec, nsub, out_s, pos = 0, 0, [], 0
    letters = "ABCDEFGH"
    for m in re.finditer(r"\\(section|subsection)\{", s):
        k = m.start(); e = arg_end(s, m.end() - 1); ttl = s[m.end():e - 1]
        if m.group(1) == "section":
            nsec += 1; nsub = 0
            if k > app:
                num = letters[nsec - 1 - n_main]
            else:
                num = str(nsec); n_main = nsec
            new = f"\\section*{{{num}\u2002{ttl}}}"
        else:
            nsub += 1
            new = f"\\subsection*{{{num}.{nsub}\u2002{ttl}}}"
        out_s.append(s[pos:k]); out_s.append(new); pos = e
    out_s.append(s[pos:]); s = "".join(out_s)
    s = s.replace("\\appendix", "")
    for lab, num in labels.items():   # every labelled section must match what we printed
        if lab.startswith("sec:"):
            assert re.search(r"\\(sub)?section\*\{" + re.escape(num) + "\u2002", s), (lab, num)

    # run-in paragraph headings -> bold lead-in
    s = replace_cmd(s, "paragraph", lambda x: "\\textbf{" + x + "} ")

    # bibliography -> plain list
    b0 = s.index("\\begin{thebibliography}"); b1 = s.index("\\end{thebibliography}") + len("\\end{thebibliography}")
    items = re.split(r"\\bibitem\[[^\]]*\]\{\w+\}", s[b0:b1])[1:]
    refs = "\\section*{References}\n\n" + "\n\n".join(x.replace("\\end{thebibliography}", "").strip() for x in items) + "\n\n"
    s = s[:b0] + refs + s[b1:]

    # abstract
    s = s.replace("\\begin{abstract}", "\\section*{Abstract}\n").replace("\\end{abstract}", "")

    # figures: PNG copies
    work = pathlib.Path(tempfile.mkdtemp())
    (work / "fig").mkdir()
    for pdf in sorted((src / "fig").glob("*.pdf")):
        subprocess.run(["pdftoppm", "-r", "200", "-png", "-singlefile", str(pdf), str(work / "fig" / pdf.stem)], check=True)
    s = re.sub(r"\\includegraphics\[[^\]]*\]\{fig/([\w]+)\.pdf\}", r"\\includegraphics[width=6.3in]{fig/\1.png}", s)

    head = ("\\documentclass{article}\n\\usepackage{amsmath,amssymb}\n\\begin{document}\n"
            f"\\section*{{{first}}}\n\n\\emph{{{sub}}}\n\nAnkit Kumar, {author}\n\n{date}\n\n\\emph{{{note}}}\n\n")
    doc = head + s + "\n\\end{document}\n"
    for bad in ["\\V{", "\\ref{", "\\cite", "%%"]:
        assert bad not in doc, bad
    (work / "paper.tex").write_text(doc)
    ref = work / "reference.docx"
    make_reference(ref)
    out.parent.mkdir(parents=True, exist_ok=True)
    r = subprocess.run(["pandoc", "paper.tex", "-f", "latex", "-t", "docx", "--reference-doc", str(ref),
                        "-o", str(out.resolve())], cwd=work, capture_output=True, text=True)
    if r.returncode:
        raise SystemExit(r.stderr)
    post_process(out)
    if r.stderr.strip():
        print(r.stderr.strip())
    shutil.copy(work / "paper.tex", out.with_suffix(".pandoc.tex"))
    return out


def make_reference(path):
    """pandoc's default reference document with Times New Roman 11 pt body text."""
    data = subprocess.run(["pandoc", "-o", str(path), "--print-default-data-file", "reference.docx"], check=True)
    import docx
    from docx.shared import Pt
    d = docx.Document(str(path))
    for name in ["Normal", "Body Text", "First Paragraph", "Compact", "Abstract", "Block Text"]:
        try:
            st = d.styles[name]
        except KeyError:
            continue
        st.font.name = "Times New Roman"; st.font.size = Pt(11)
    for name, size in [("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 11), ("Title", 16)]:
        try:
            st = d.styles[name]; st.font.name = "Times New Roman"; st.font.size = Pt(size)
            st.font.color.rgb = None
        except KeyError:
            pass
    try:
        d.styles["Image Caption"].font.size = Pt(9.5)
        d.styles["Table Caption"].font.size = Pt(9.5)
    except KeyError:
        pass
    d.save(str(path))


def post_process(path):
    """Borders and small type for every table, so wide tables fit the page."""
    import docx
    from docx.shared import Pt
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    d = docx.Document(str(path))
    for sec in d.sections:
        sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = docx.shared.Cm(2.2)
    for t in d.tables:
        tblPr = t._tbl.tblPr
        b = OxmlElement("w:tblBorders")
        for edge in ("top", "bottom", "insideH"):
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "4"); el.set(qn("w:color"), "808080")
            b.append(el)
        tblPr.append(b)
        # column widths from the longest entry in each column, so no number or name wraps
        grid = t._tbl.tblGrid
        ncols = len(grid.findall(qn("w:gridCol")))
        longest = [3] * ncols; word = [3] * ncols
        for tr in t._tbl.findall(qn("w:tr")):
            c = 0
            for tc in tr.findall(qn("w:tc")):
                span = tc.find(qn("w:tcPr") + "/" + qn("w:gridSpan"))
                n = int(span.get(qn("w:val"))) if span is not None else 1
                txt = "".join(x.text or "" for x in tc.iter(qn("w:t")))
                txt += "".join(x.text or "" for x in tc.iter("{http://schemas.openxmlformats.org/officeDocument/2006/math}t"))
                if n == 1 and c < ncols:
                    longest[c] = max(longest[c], len(txt))
                    # numbers and intervals never break; words may
                    numeric = re.fullmatch(r"[\s\d.,+\-\u2212\[\]<>\u00d7%()/]*", txt) is not None
                    word[c] = max([word[c], len(txt) if numeric else 0] + [len(x) for x in txt.split()])
                c += n
        # width each column needs for its longest entry on one line, and the least it
        # can have (its longest word); text columns wrap first, numbers never
        total = 9400
        pt_ = 7.5 if ncols >= 9 else 8 if ncols >= 7 else 9
        cw, pad = pt_ * 0.52 * 20, 240
        need = [x * cw + pad for x in longest]; least = [x * cw + pad for x in word]
        if sum(need) <= total:
            widths = [int(x + (total - sum(need)) * x / sum(need)) for x in need]
        else:
            slack = [a - b for a, b in zip(need, least)]; deficit = sum(need) - total
            if sum(slack) >= deficit:
                widths = [int(a - deficit * s / sum(slack)) for a, s in zip(need, slack)]
            else:
                widths = [int(total * x / sum(least)) for x in least]
        for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
            gc.set(qn("w:w"), str(w))
        for el in tblPr.findall(qn("w:tblW")) + tblPr.findall(qn("w:tblLayout")):
            tblPr.remove(el)
        w = OxmlElement("w:tblW"); w.set(qn("w:w"), str(sum(widths))); w.set(qn("w:type"), "dxa"); tblPr.append(w)
        lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
        for tr in t._tbl.findall(qn("w:tr")):
            c = 0
            for tc in tr.findall(qn("w:tc")):
                span = tc.find(qn("w:tcPr") + "/" + qn("w:gridSpan"))
                n = int(span.get(qn("w:val"))) if span is not None else 1
                tcw = tc.find(qn("w:tcPr") + "/" + qn("w:tcW"))
                if tcw is not None:
                    tcw.set(qn("w:type"), "dxa"); tcw.set(qn("w:w"), str(sum(widths[c:c + n])))
                c += n
        ncol = len(t.columns)
        size = Pt(7.5) if ncol >= 9 else Pt(8) if ncol >= 7 else Pt(9)
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(0)
                    for run in p.runs:
                        run.font.size = size
    d.save(str(path))


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--src", default=str(HERE / "out" / "paper2"))
    ap.add_argument("--out", default=str(HERE.parent / "paper" / "two_backbones" / "KernelRoute_paper.docx"))
    a = ap.parse_args()
    print(build(a.src, a.out))
